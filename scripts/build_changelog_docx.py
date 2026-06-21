r"""
Build the v1.2 change-log addendum as a formatted Word document.

Generates a standalone .docx (it does NOT modify the existing final-submission
PDFs/DOCX) summarising the alerting/notification and popularity changes, so it
can sit alongside the other deliverables in the Final Submission folder.

Usage:
    .\venv\Scripts\python.exe -m scripts.build_changelog_docx
"""
from __future__ import annotations

import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = Path(r"C:\Users\ymk00\Desktop\Final Subbmission\Change_Log_Addendum_v1.2.docx")

ACCENT = RGBColor(0x1F, 0x4E, 0x79)


def main() -> int:
    doc = Document()

    # --- base styles ---
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    # --- title ---
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Smart Reception Assistant")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = ACCENT

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = sub.add_run("Change Log Addendum — v1.2")
    r.bold = True
    r.font.size = Pt(14)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.add_run("Date: 2026-06-20    |    Scope: Alerting/Notifications (FR-4), "
                 "Upselling Popularity (FR-2), Test Tooling").italic = True

    doc.add_paragraph(
        "This addendum documents changes made after the v1.1 baseline. It "
        "supplements the SDD and SysRS rather than replacing them. No changes "
        "were made to the face-recognition pipeline."
    )

    def h(text: str, level: int = 1):
        p = doc.add_heading(text, level=level)
        for rn in p.runs:
            rn.font.color.rgb = ACCENT
        return p

    def bullets(items):
        for it in items:
            doc.add_paragraph(it, style="List Bullet")

    def table(headers, rows):
        t = doc.add_table(rows=1, cols=len(headers))
        t.style = "Light Grid Accent 1"
        for i, head in enumerate(headers):
            cell = t.rows[0].cells[i]
            cell.text = ""
            run = cell.paragraphs[0].add_run(head)
            run.bold = True
        for row in rows:
            cells = t.add_row().cells
            for i, val in enumerate(row):
                cells[i].text = str(val)
        doc.add_paragraph()

    # --- 1. Notification templates ---
    h("1. Per-Incident Notification Templates (FR-4)")
    doc.add_paragraph(
        "Previously every alert e-mail used one generic plain-text body, so all "
        "incident types looked identical and offered no guidance. A new module, "
        "modules/alerts/notification_templates.py, now renders a tailored "
        "notification for each AlertType. A single _INCIDENTS table drives the "
        "presentation, so each incident keeps one identity (icon, label, accent "
        "colour, and a clear call to action) across every channel."
    )
    table(
        ["Incident", "Subject prefix", "Call to action"],
        [
            ["security", "Security alert", "Dispatch a staff member to identify this person."],
            ["wanted", "Watchlist match", "Notify security immediately; do not approach."],
            ["assistance", "Guest needs assistance", "Check in with the waiting guest."],
            ["vip_arrival", "VIP arrival", "Prepare the VIP welcome; notify the duty manager."],
            ["arrival", "Guest arrival", "Greet the arriving guest and begin check-in."],
            ["due_out", "Guest due out", "A due-out guest is standing at the reception."],
            ["(unknown)", "Reception alert", "Review this alert in the dashboard."],
        ],
    )
    doc.add_paragraph("Public API:")
    bullets([
        "render_alert_email(alert) -> (subject, text_body, html_body): branded, "
        "severity-coloured HTML card with a plain-text fallback (sent multipart).",
        "render_whatsapp(alert) -> str: compact message using WhatsApp *bold* / "
        "_italic_ markup (degrades cleanly to plain SMS).",
        "Zero new dependencies; fail-soft fallback to a generic block for unknown "
        "alert types.",
    ])

    # --- 2. WhatsApp / Twilio ---
    h("2. WhatsApp / Twilio Alert Channel (FR-4)")
    doc.add_paragraph(
        "The SMS path was extended to support Twilio's WhatsApp Sandbox in "
        "addition to plain SMS."
    )
    bullets([
        "New setting twilio_channel (sms | whatsapp), read from env TWILIO_CHANNEL.",
        "send_sms() adds the 'whatsapp:' prefix to the From and each recipient when "
        "the channel is whatsapp (helper _wa_address); SMS behaviour is unchanged.",
        "The message body reuses render_whatsapp() — the same per-incident identity "
        "as e-mail.",
        "twilio==9.3.0 installed in the venv.",
    ])
    doc.add_paragraph(
        "Operational requirement: each recipient must join the sandbox "
        "('join <code>' to +1 415 523 8886) and, on a trial account, be a Verified "
        "Caller ID. The 72-hour sandbox session must be refreshed periodically."
    )

    # --- 3. BCC ---
    h("3. Recipient Privacy — Blind Copy (BCC)")
    doc.add_paragraph(
        "Alert recipients were previously listed in the To header, exposing every "
        "notified person's address to the others. Recipients are now placed in Bcc. "
        "A cosmetic To header (the sending mailbox) avoids an 'undisclosed-recipients' "
        "display, and the delivery envelope is passed explicitly so the Bcc header is "
        "stripped before transmission and only the listed recipients receive the mail. "
        "Verified at the wire level: no Bcc header is transmitted and the envelope "
        "equals the recipient list."
    )

    # --- 4. Popularity ---
    h("4. Statistics-Driven Service Popularity (FR-2)")
    doc.add_paragraph(
        "Service.popularity_score was a hand-picked constant. The server now derives "
        "it from recommendation outcomes using a Laplace-smoothed acceptance rate "
        "(Beta(1,1) prior): score = (accepted + 1) / (accepted + declined + 2)."
    )
    bullets([
        "No history -> neutral 0.5; converges to accepted / (accepted + declined) "
        "with volume.",
        "Smoothing keeps a service shown once/accepted once (0.667) from outranking "
        "one shown 100x/accepted 80x (0.794).",
        "ServiceRepository.recompute_popularity(db, service_id=None) recomputes one "
        "service or the whole catalogue.",
        "Live update: a service's popularity refreshes when a guest's recommendation "
        "is marked ACCEPTED or DECLINED.",
        "Admin endpoint: POST /api/services/recompute-popularity (admin/manager).",
    ])
    doc.add_paragraph(
        "The recommendation engine still reads popularity_score as its R1 baseline; "
        "only the source of the number changed."
    )

    # --- 5. Mock data ---
    h("5. Mock Test Population — 100 Guests")
    doc.add_paragraph(
        "scripts/seed_mock_guests.py populates the database for full-system testing."
    )
    bullets([
        "100 realistic guests: varied titles; 12 nationality/language locales; ~15% "
        "VIP; dietary/room/pillow preferences; mixed statuses; ~4% watch-listed for "
        "security testing.",
        "~400 recommendation accept/decline interactions (per-service hidden 'appeal') "
        "so the popularity statistics are meaningful, followed by a real "
        "recompute_popularity pass.",
        "Face embeddings are deliberately NOT mocked — the recognition model is only "
        "meaningful against real faces; these guests are profile-registered but not "
        "face-enrolled.",
        "Idempotent (@mock.example.com domain); flags --count and --reset-interactions.",
    ])

    # --- 6. HTTPS/WSS ---
    h("6. HTTPS / WSS Transport via Caddy (NFR-2.2 / NFR-2.3)")
    doc.add_paragraph(
        "The dashboard, REST API and /ws WebSocket were previously served over "
        "plain HTTP/WS. A Caddyfile in the repository root now puts a Caddy reverse "
        "proxy in front of the FastAPI app: Caddy terminates TLS while the app keeps "
        "speaking plain HTTP/WS on 127.0.0.1:5000, and /ws upgrades to WSS "
        "automatically. No application code changed."
    )
    bullets([
        "localhost site block: local/demo with an auto-trusted certificate.",
        "LAN IP with 'tls internal': run 'caddy trust' once per kiosk for "
        "warning-free HTTPS across the network.",
        "Public domain: automatic, auto-renewing Let's Encrypt certificate.",
        "Run with 'caddy run' from the repo root; generate_self_signed_cert.py "
        "remains as a direct-to-uvicorn alternative.",
        "Closes NFR-2.2 (HTTPS) and NFR-2.3 (WSS). Caddyfile validated with "
        "'caddy validate' (Caddy v2.11.4): Valid configuration.",
    ])

    # --- 7. Config ---
    h("7. Configuration Fixes & Additions")
    bullets([
        "Latent bug fixed: config.yaml keyed the block 'alerts:' but the loader "
        "reads 'alert:', so the whole alert section was silently ignored and defaults "
        "were used. Renamed to 'alert:' so configured values actually load.",
        "sms_enabled set to true in config.yaml.",
        "New env var TWILIO_CHANNEL documented in .env.example.",
    ])

    # --- 8. Files ---
    h("8. New / Modified Files")
    doc.add_paragraph("New files:").runs[0].bold = True
    table(
        ["File", "Purpose"],
        [
            ["modules/alerts/notification_templates.py", "Per-incident email + WhatsApp templates"],
            ["scripts/seed_mock_guests.py", "100-guest mock population + interaction history"],
            ["scripts/send_test_whatsapp.py", "Manual Twilio WhatsApp/SMS send test"],
            ["Caddyfile", "Caddy reverse proxy -> HTTPS/WSS (NFR-2.2 / NFR-2.3)"],
        ],
    )
    doc.add_paragraph("Modified files:").runs[0].bold = True
    table(
        ["File", "Change"],
        [
            ["modules/alerts/alert_notifier.py", "Use templates; WhatsApp channel; BCC privacy"],
            ["config/settings.py", "twilio_channel setting"],
            ["config.yaml", "alerts->alert key fix; sms_enabled: true"],
            [".env / .env.example", "Twilio WhatsApp configuration"],
            ["database/repositories/service_repository.py", "popularity_from_stats, recompute_popularity"],
            ["api/routes/recommendations.py", "Live popularity recompute on accept/decline"],
            ["api/routes/services.py", "POST /recompute-popularity endpoint"],
        ],
    )

    # --- 9. Testing ---
    h("9. Testing & Verification")
    bullets([
        "pytest tests/test_recommendations.py tests/test_upselling.py — 42 passed.",
        "Notification templates render for all incident types (email + WhatsApp).",
        "BCC verified at the SMTP wire level: Bcc header absent on transmission; "
        "envelope equals the recipient list.",
        "seed_mock_guests.py executed: 100 guests + 402 interactions; popularity "
        "recomputed (0.19 – 0.82); re-run is idempotent.",
        "HTTPS/WSS: Caddyfile validated with 'caddy validate' (Caddy v2.11.4) — "
        "Valid configuration; a live end-to-end browser session is not yet exercised.",
    ])

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Wrote {OUT}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
